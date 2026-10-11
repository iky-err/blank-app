import io
import os
import time
import uuid
from datetime import datetime

import pandas as pd
import streamlit as st
from docx import Document
from docx.shared import Pt, RGBColor

from claude_client import call_claude_json
from prompts import (
    SYSTEM_SCAN, SYSTEM_PATHS, SYSTEM_EBIOS,
    SOURCES_RISQUE, OBJECTIFS_VISES, NIVEAUX_GRAVITE, NIVEAUX_VRAISEMBLANCE, MAX_SCORE,
    SCENARIOS_REFERENCE, INHERITED_FIELDS,
    format_refs, format_paths, score_scenario, validate_path, validate_scenario,
    warnings_path, warnings_scenario,
)

st.set_page_config(page_title="Formind — Outil EBIOS RM", page_icon="🛡️", layout="wide")

SECTEURS = [
    "Finance / Banque", "Santé", "Énergie / OIV", "Collectivité", "Industrie",
    "Tech / SaaS", "Transport", "Défense", "Assurance", "Retail", "Autre",
]
TAILLES = [
    "TPE (<50 sal.)", "PME (50-500 sal.)", "ETI (500-5000 sal.)",
    "GE (>5000 sal.)", "Administration", "OIV / OSE",
]
REGLEMENTS = ["RGPD", "NIS2", "PCI-DSS", "ISO 27001", "DORA", "SOC 2", "Loi 09-08 (CNDP)", "Loi 05-20 (DGSSI)"]
MAT_DIMS = [
    ("gov", "Gouvernance & Organisation"),
    ("risk", "Gestion des risques"),
    ("tech", "Protection technique"),
    ("inc", "Gestion des incidents"),
    ("conf", "Conformité & Audit"),
]
MAT_LABELS = {1: "Initial", 2: "Répété", 3: "Défini", 4: "Géré", 5: "Optimisé"}
NB_CHEMINS_OPTIONS = [5, 7, 10, 12, 15]
BATCH_SIZE = 4  # chemins traités par appel, pour éviter les réponses JSON tronquées
STATUS_LABELS = {"pending": "🟡 En attente", "validated": "🟢 Validé", "rejected": "🔴 Rejeté"}
AUTRE = "Autre (saisie libre)…"
MOTIFS_REJET = [
    "Peu plausible / irréaliste pour ce client", "Hors périmètre de l'étude", "Doublon d'une autre proposition",
    "Incohérent avec la méthode EBIOS RM", "Trop générique / pas assez précis", "Autre",
]

# Champs que le consultant peut corriger : toute différence avec la proposition IA = « modifié ».
PATH_EDITABLE = [
    "source_risque", "objectif_vise", "partie_prenante", "valeur_metier",
    "evenement_redoute", "gravite", "chemin", "ref_type",
]
SCEN_EDITABLE = [
    "titre", "mode_operatoire", "actif", "vuln", "impact_op", "impact_fin",
    "impact_reg", "vraisemblance", "reco",
]
# Indicateurs : enregistrés en arrière-plan, jamais affichés au consultant ni exportés au client.
# Une ligne par étude, mise à jour à chaque action. Emplacement modifiable via la variable d'environnement.
INDIC_FILE = os.environ.get(
    "EBIOS_INDICATEURS", os.path.join(os.path.dirname(os.path.abspath(__file__)), "indicateurs_ebios.csv"))

# Exemple FICTIF de cadrage (aucun client réel), téléchargeable en Word comme modèle de saisie.
# (libellé du champ dans l'outil, exemple de réponse, conseil de saisie)
EXEMPLE = [
    ("Nom de l'organisation *", "PayNova (exemple fictif)", "Nom du client tel qu'il apparaîtra dans l'export."),
    ("Taille", "PME (50-500 sal.)", "Choisir dans la liste."),
    ("Secteur d'activité", "Finance / Banque", "Choisir « Autre » et préciser si le secteur n'est pas listé."),
    ("Périmètre de l'étude (métier et technique) *",
     "Activité d'émission de paiements mobiles et SI Finance associé, siège de Casablanca.",
     "Indiquer l'activité étudiée ET le SI concerné, avec les sites si utile."),
    ("Valeurs métier",
     "Exécution des paiements clients ; gestion des comptes de paiement ; reporting réglementaire à Bank Al-Maghrib.",
     "Processus ou informations essentiels, séparés par des points-virgules."),
    ("Biens supports",
     "Plateforme de paiement (cloud) → exécution des paiements ; base clients → gestion des comptes ; "
     "ERP Finance → reporting ; messagerie Microsoft 365 ; équipe trésorerie ; prestataire d'infogérance.",
     "Format conseillé « bien support → valeur métier soutenue »."),
    ("Événements redoutés (avec gravité si connue)",
     "Détournement de fonds — Critique ; indisponibilité de la plateforme > 24 h — Grave ; fuite de données clients — Grave.",
     "Ajouter la gravité (Mineure, Significative, Grave, Critique) : elle sera reprise dans les chemins."),
    ("Sources de risque pressenties", "Cybercriminalité organisée ; Prestataire ou partenaire malveillant",
     "Facultatif : laisser vide pour que l'IA propose elle-même les sources plausibles."),
    ("Scénarios de référence à couvrir",
     "Erreur de configuration (Finance) ; Usurpation d'identité à des fins de fraude ; Sabotage contractuel",
     "Chaque scénario coché donnera au moins un chemin d'attaque."),
    ("Référentiels et réglementations applicables", "Loi 09-08 (CNDP) ; PCI-DSS ; ISO 27001", "Cocher tous ceux qui s'appliquent."),
    ("Effectif équipe IT / équipe sécurité", "12 / 2", "Nombre de personnes (internes et externes dédiées)."),
    ("Parties prenantes de l'écosystème",
     "Infogérant (administration des serveurs), éditeur de l'ERP, banque partenaire, agrégateur SMS.",
     "Préciser leur rôle ou leurs accès : l'IA s'en sert comme vecteurs d'attaque."),
    ("Projets SI en cours ou récents", "Migration de la plateforme vers le cloud (en cours).", ""),
    ("Incidents de sécurité passés",
     "Tentative de fraude au président par e-mail en 2025, détectée par la trésorerie.", "Écrire « Aucun » s'il n'y en a pas."),
    ("État de la menace connu", "Hausse des campagnes de phishing ciblant les fintechs de la région.",
     "Veille, alertes sectorielles, retours du client."),
    ("Maturité SSI (5 curseurs, 1 à 5)", "Gouvernance 2 ; Risques 2 ; Protection technique 3 ; Incidents 2 ; Conformité 3",
     "1 = Initial, 3 = Défini, 5 = Optimisé."),
    ("Nombre de chemins d'attaque à proposer", "7", "Entre 5 et 15 ; au moins un par scénario de référence."),
]


def exemple_word() -> bytes:
    """Fiche d'exemple de cadrage au format Word (modèle de saisie pour le consultant)."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Outil EBIOS RM — Exemple de cadrage (Atelier 1)", level=1)
    p = doc.add_paragraph(
        "Cet exemple est FICTIF. Il montre le niveau de détail attendu pour chaque champ de l'étape 1. "
        "Saisissez dans l'outil les informations réelles de votre client : l'IA n'analyse que ce que vous saisissez."
    )
    p.runs[0].italic = True
    doc.add_paragraph("Les champs marqués * sont obligatoires ; les autres améliorent fortement la qualité des propositions.")
    t = doc.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    for cell, txt in zip(t.rows[0].cells, ("Champ de l'outil", "Exemple de réponse", "Conseil de saisie")):
        cell.text = ""
        cell.paragraphs[0].add_run(txt).bold = True
    for champ, ex, conseil in EXEMPLE:
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = champ, ex, conseil
        row[0].paragraphs[0].runs[0].bold = True
        if row[2].paragraphs[0].runs:
            row[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _new_metrics() -> dict:
    return {
        "id_etude": uuid.uuid4().hex[:8], "date_debut": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "t_debut": time.time(), "duree_premiere_liste": None, "duree_premiers_scenarios": None,
        "temps_ia": 0.0, "nb_appels_ia": 0, "nb_generations_chemins": 0,
        "refs_manquees_ia": [], "chrono_reinitialise": False,
    }


defaults = {
    "step": 1,
    "nom": "", "secteur": SECTEURS[0], "secteur_autre": "",
    "taille": TAILLES[1], "perimetre": "", "biens_supports": "",
    "valeurs_metier": "", "evenements_redoutes": "", "sources_risque": [],
    "scenarios_ref": ["REF-CONFIG", "REF-USURP", "REF-SABOT"], "scenarios_ref_autre": "",
    "reglements": [], "reglement_autre": "",
    "incidents": "", "menaces": "",
    "equipe_it": "", "equipe_sec": "", "prestataires": "", "projets_recents": "",
    "notes": "",
    "maturite": {k: 2 for k, _ in MAT_DIMS},
    "nb_chemins": 7,
    "analysis": None,
    "paths": None,
    "path_status": {},
    "scenarios": None,
    "scen_status": {},
    "flash": "",
}
for k, v in defaults.items():
    st.session_state.setdefault(k, v)
st.session_state.setdefault("metrics", _new_metrics())


# ── Helpers ────────────────────────────────────────────────────────────────

def _pipe(v) -> str:
    """Forme canonique d'une liste « a | b | c » (sert à comparer avant / après correction)."""
    return " | ".join(x.strip() for x in str(v or "").replace("\n", "|").split("|") if x.strip())


def _norm(v) -> str:
    return " ".join(str(v or "").split())


def is_modified(item: dict, fields: list[str]) -> bool:
    orig = item.get("_orig")
    if not orig:
        return False
    return any(_norm(item.get(f)) != _norm(orig.get(f)) for f in fields)


def fmt_duree(sec) -> str:
    if sec is None:
        return "—"
    sec = int(round(sec))
    m, s = divmod(sec, 60)
    h, m = divmod(m, 60)
    return f"{h} h {m:02d} min" if h else (f"{m} min {s:02d} s" if m else f"{s} s")


def reset_state() -> None:
    """Vide la session (y compris les widgets) pour démarrer une nouvelle étude."""
    for k in list(st.session_state.keys()):
        del st.session_state[k]
    for k, v in defaults.items():
        st.session_state[k] = v
    st.session_state["metrics"] = _new_metrics()


def call_ia(system: str, user: str):
    """Appel au modèle, chronométré pour les indicateurs."""
    met = st.session_state["metrics"]
    t = time.time()
    try:
        return call_claude_json(system, user)
    finally:
        met["temps_ia"] += time.time() - t
        met["nb_appels_ia"] += 1


def erreur_ia(e: Exception) -> None:
    st.error(
        f"La génération n'a pas abouti : {e}\n\n"
        "Vos saisies sont conservées. Relancez le bouton ; si l'erreur persiste, vérifiez la connexion "
        "et la clé API (fichier de configuration)."
    )


def get_refs() -> list[dict]:
    """Scénarios de référence du consultant : catalogue coché + lignes libres."""
    m = st.session_state
    refs = [{"id": rid, **SCENARIOS_REFERENCE[rid]} for rid in m["scenarios_ref"] if rid in SCENARIOS_REFERENCE]
    for i, line in enumerate(m["scenarios_ref_autre"].splitlines(), start=1):
        if line.strip():
            refs.append({"id": f"REF-C{i:02d}", "label": line.strip(), "desc": ""})
    return refs


def build_context() -> str:
    m = st.session_state
    mat_str = ", ".join(
        f"{label}: {m['maturite'][k]}/5 ({MAT_LABELS[m['maturite'][k]]})"
        for k, label in MAT_DIMS
    )
    regs = list(m["reglements"]) + ([m["reglement_autre"]] if m["reglement_autre"] else [])
    return f"""Organisation: {m['nom']}
Secteur: {m['secteur_autre'] if m['secteur'] == 'Autre' else m['secteur']} | Taille: {m['taille']}
Périmètre de l'étude (métier et technique): {m['perimetre'] or 'Non précisé'}
Valeurs métier: {m['valeurs_metier'] or 'Non précisées'}
Biens supports: {m['biens_supports'] or 'Non précisés'}
Événements redoutés identifiés: {m['evenements_redoutes'] or 'Non précisés'}
Sources de risque pressenties (consultant): {', '.join(m['sources_risque']) or 'Non précisées'}
Réglementation: {', '.join(regs) or 'Non précisée'}
Incidents passés: {m['incidents'] or 'Aucun signalé'}
État de la menace connu: {m['menaces'] or 'Non précisé'}
Équipe IT: {m['equipe_it'] or '?'} | Équipe sécurité: {m['equipe_sec'] or '?'}
Parties prenantes de l'écosystème: {m['prestataires'] or '?'}
Projets récents: {m['projets_recents'] or '?'}
Maturité SSI (indicateur complémentaire, hors EBIOS RM): {mat_str}
Notes consultant: {m['notes'] or 'Aucune'}"""


def choix_ou_saisie(container, label: str, options: list[str], current: str, key: str) -> str:
    """Liste déroulante avec une option « Autre » qui ouvre un champ de saisie libre."""
    opts = options + [AUTRE]
    hors_liste = current not in options
    choix = container.selectbox(label, opts, index=opts.index(AUTRE) if hors_liste else opts.index(current), key=f"{key}_sel")
    if choix != AUTRE:
        return choix
    libre = container.text_input(f"{label} — précisez", value=current if hors_liste else "", key=f"{key}_txt")
    if not libre.strip():
        container.caption("Saisissez une valeur : en attendant, la valeur précédente est conservée.")
        return current
    return libre.strip()


def validated_paths() -> list[dict]:
    m = st.session_state
    return [p for p in (m["paths"] or []) if m["path_status"].get(p["id"]) == "validated"]


def compute_coverage() -> list[dict]:
    """Pour chaque scénario de référence : chemin proposé ? validé ? scénario généré ?"""
    m = st.session_state
    paths = m["paths"] or []
    scenarios = [s for s in (m["scenarios"] or []) if m["scen_status"].get(s["id"]) != "rejected"]
    rows = []
    for r in get_refs():
        r_paths = [p for p in paths if p.get("ref_type") == r["id"]]
        r_valid = [p for p in r_paths if m["path_status"].get(p["id"]) == "validated"]
        r_scen = [s for s in scenarios if s.get("ref_type") == r["id"]]
        if r_scen:
            state = "✅ Couvert"
        elif r_valid:
            state = "⚠️ Chemin validé, scénario à générer"
        elif r_paths:
            state = "🟡 Chemin proposé, non validé"
        else:
            state = "❌ Aucun chemin d'attaque"
        rows.append({
            "Scénario de référence": r["label"],
            "Chemins": ", ".join(p["id"] for p in r_paths) or "—",
            "Scénarios": ", ".join(s["id"] for s in r_scen) or "—",
            "Couverture": state,
        })
    return rows


# ── Indicateurs ────────────────────────────────────────────────────────────

def _decisions(items: list[dict], status: dict, fields: list[str]) -> dict:
    n = len(items)
    tel_quel = sum(1 for i in items if status.get(i["id"]) == "validated" and not is_modified(i, fields))
    modifie = sum(1 for i in items if status.get(i["id"]) == "validated" and is_modified(i, fields))
    rejete = sum(1 for i in items if status.get(i["id"]) == "rejected")
    pct = lambda x: round(100 * x / n, 1) if n else 0.0
    return {
        "n": n, "tel_quel": tel_quel, "modifie": modifie, "rejete": rejete,
        "attente": n - tel_quel - modifie - rejete,
        "pct_tel_quel": pct(tel_quel), "pct_modifie": pct(modifie), "pct_rejete": pct(rejete),
        "pct_attente": pct(n - tel_quel - modifie - rejete),
    }


def compute_metrics() -> dict:
    m = st.session_state
    met = m["metrics"]
    paths = m["paths"] or []
    ia_paths = [p for p in paths if p.get("_origine", "ia") == "ia"]
    manuels = [p for p in paths if p.get("_origine") == "manuel"]
    manuels_importants = [p for p in manuels if p["gravite"] in ("Grave", "Critique")]
    scen = m["scenarios"] or []
    motifs = {}
    for it in ia_paths + scen:
        if m["path_status"].get(it["id"]) == "rejected" or m["scen_status"].get(it["id"]) == "rejected":
            k = it.get("_motif_rejet") or "Non précisé"
            motifs[k] = motifs.get(k, 0) + 1
    return {
        "chemins": _decisions(ia_paths, m["path_status"], PATH_EDITABLE),
        "scenarios": _decisions(scen, m["scen_status"], SCEN_EDITABLE),
        "manuels": manuels, "manuels_importants": manuels_importants,
        "refs_manquees_ia": met.get("refs_manquees_ia", []),
        "motifs_rejet": motifs,
        "duree_premiere_liste": met.get("duree_premiere_liste"),
        "duree_premiers_scenarios": met.get("duree_premiers_scenarios"),
        "temps_ia": met.get("temps_ia", 0.0), "nb_appels_ia": met.get("nb_appels_ia", 0),
        "nb_generations_chemins": met.get("nb_generations_chemins", 0),
        "chrono_reinitialise": met.get("chrono_reinitialise", False),
    }


def save_indicators() -> None:
    """Écrit (ou met à jour) la ligne de l'étude dans le fichier d'indicateurs. Silencieux en cas d'erreur."""
    m = st.session_state
    try:
        k = compute_metrics()
        met = m["metrics"]
        mn = lambda sec: round(sec / 60, 1) if sec is not None else None
        row = {
            "id_etude": met.get("id_etude"), "date_debut": met.get("date_debut"),
            "organisation": m["nom"], "secteur": m["secteur_autre"] if m["secteur"] == "Autre" else m["secteur"],
            "duree_1re_liste_min": mn(k["duree_premiere_liste"]),
            "duree_1ers_scenarios_min": mn(k["duree_premiers_scenarios"]),
            "temps_ia_min": mn(k["temps_ia"]), "nb_appels_ia": k["nb_appels_ia"],
            "nb_generations_chemins": k["nb_generations_chemins"],
        }
        for pref, d in (("chemins", k["chemins"]), ("scenarios", k["scenarios"])):
            row.update({
                f"{pref}_proposes_ia": d["n"], f"{pref}_tel_quel": d["tel_quel"],
                f"{pref}_modifies": d["modifie"], f"{pref}_rejetes": d["rejete"], f"{pref}_en_attente": d["attente"],
                f"{pref}_pct_tel_quel": d["pct_tel_quel"], f"{pref}_pct_modifies": d["pct_modifie"],
                f"{pref}_pct_rejetes": d["pct_rejete"],
            })
        row.update({
            "manques_ia": len(k["manuels"]), "manques_ia_importants": len(k["manuels_importants"]),
            "manques_ia_detail": " ; ".join(
                f"{p['evenement_redoute']} ({p['gravite']})" + (f" — {p['_motif_ajout']}" if p.get("_motif_ajout") else "")
                for p in k["manuels"]),
            "refs_non_couvertes_ia": " ; ".join(k["refs_manquees_ia"]),
            "motifs_rejet": " ; ".join(f"{a} : {b}" for a, b in k["motifs_rejet"].items()),
            "chrono_reinitialise": k["chrono_reinitialise"],
            "etude_complete": bool(k["scenarios"]["n"]) and k["chemins"]["attente"] == 0 and k["scenarios"]["attente"] == 0,
        })
        if row == m.get("_derniere_ligne"):
            return  # rien n'a changé depuis la dernière écriture
        m["_derniere_ligne"] = dict(row)
        row["derniere_maj"] = datetime.now().strftime("%Y-%m-%d %H:%M")
        df = pd.read_csv(INDIC_FILE, dtype={"id_etude": str}) if os.path.exists(INDIC_FILE) else pd.DataFrame()
        if not df.empty:
            df = df[df["id_etude"] != row["id_etude"]]
        df = pd.concat([df, pd.DataFrame([row])], ignore_index=True)
        df.to_csv(INDIC_FILE, index=False, encoding="utf-8-sig")
    except Exception:
        pass  # les indicateurs ne doivent jamais bloquer le consultant


# ── Génération ─────────────────────────────────────────────────────────────

def generate_paths() -> None:
    m = st.session_state
    met = m["metrics"]
    refs = get_refs()
    n = max(m["nb_chemins"], len(refs))
    raw = call_ia(
        SYSTEM_PATHS,
        f"Propose {n} chemins d'attaque stratégiques (au minimum un par scénario de référence).\n\n"
        f"CONTEXTE:\n{build_context()}\n\n"
        f"SCÉNARIOS DE RÉFÉRENCE À COUVRIR:\n{format_refs(refs) or 'Aucun'}",
    )
    if not isinstance(raw, list):
        raise ValueError("Format de réponse inattendu (tableau attendu).")

    valid_refs = {r["id"] for r in refs}
    paths = []
    for item in raw:
        p = validate_path(item, valid_refs)
        if p is None:
            continue
        p["id"] = f"CA-{len(paths) + 1:02d}"
        p["_origine"] = "ia"
        p["_orig"] = {f: p.get(f, "") for f in PATH_EDITABLE}
        paths.append(p)
    if not paths:
        raise ValueError("Aucun chemin d'attaque exploitable n'a pu être extrait. Réessayez.")
    if len(paths) < len(raw):
        m["flash"] = f"{len(raw) - len(paths)} chemin(s) incomplet(s) ignoré(s)."

    m["paths"] = paths
    m["path_status"] = {p["id"]: "pending" for p in paths}
    m["scenarios"] = None
    m["scen_status"] = {}

    met["nb_generations_chemins"] += 1
    if met["duree_premiere_liste"] is None:
        met["duree_premiere_liste"] = time.time() - met["t_debut"]
    couverts = {p["ref_type"] for p in paths if p.get("ref_type")}
    met["refs_manquees_ia"] = [r["label"] for r in refs if r["id"] not in couverts]


def inherit(s: dict, path: dict) -> dict:
    """Recopie les champs du chemin validé dans le scénario et recalcule le score."""
    for f in INHERITED_FIELDS:
        s[f] = path.get(f, "")
    s["score"] = score_scenario(s["vraisemblance"], s["gravite"])
    return s


def generate_scenarios_for(paths_todo: list[dict], start: int) -> list[dict]:
    refs = get_refs()
    by_id = {p["id"]: p for p in paths_todo}
    out, n = [], start
    for i in range(0, len(paths_todo), BATCH_SIZE):
        chunk = paths_todo[i:i + BATCH_SIZE]
        chunk_ids = {p["id"] for p in chunk}
        raw = call_ia(
            SYSTEM_EBIOS,
            "Génère un scénario opérationnel pour chacun des chemins d'attaque validés ci-dessous.\n\n"
            f"CONTEXTE:\n{build_context()}\n\n"
            f"SCÉNARIOS DE RÉFÉRENCE:\n{format_refs(refs) or 'Aucun'}\n\n"
            f"CHEMINS D'ATTAQUE VALIDÉS:\n{format_paths(chunk)}",
        )
        if not isinstance(raw, list):
            raise ValueError("Format de réponse inattendu (tableau attendu).")
        seen = set()
        for item in raw:
            s, err = validate_scenario(item, chunk_ids)
            if err or s["chemin_id"] in seen:
                continue
            seen.add(s["chemin_id"])
            s["mode_operatoire"], s["reco"] = _pipe(s["mode_operatoire"]), _pipe(s["reco"])
            s = inherit(s, by_id[s["chemin_id"]])
            s["id"] = f"SO-{n:03d}"
            s["_orig"] = {f: s.get(f, "") for f in SCEN_EDITABLE}
            n += 1
            out.append(s)
    return out


def build_scenarios() -> None:
    """Conserve les scénarios des chemins toujours validés, génère ceux qui manquent."""
    m = st.session_state
    met = m["metrics"]
    valid = validated_paths()
    by_id = {p["id"]: p for p in valid}

    kept = [inherit(s, by_id[s["chemin_id"]]) for s in (m["scenarios"] or []) if s["chemin_id"] in by_id]
    have = {s["chemin_id"] for s in kept}
    todo = [p for p in valid if p["id"] not in have]

    start = max((int(s["id"].split("-")[1]) for s in kept), default=0) + 1
    new = generate_scenarios_for(todo, start) if todo else []

    scenarios = kept + new
    if not scenarios:
        raise ValueError("Aucun scénario exploitable n'a pu être extrait. Réessayez la génération.")

    missing = len(todo) - len(new)
    if missing:
        m["flash"] = f"{missing} chemin(s) validé(s) sans scénario exploitable : relancez la génération des scénarios manquants."

    m["scenarios"] = scenarios
    m["scen_status"] = {s["id"]: m["scen_status"].get(s["id"], "pending") for s in scenarios}
    if new and met["duree_premiers_scenarios"] is None:
        met["duree_premiers_scenarios"] = time.time() - met["t_debut"]


# ── Export ─────────────────────────────────────────────────────────────────

def export_excel(scenarios: list[dict]) -> bytes:
    m = st.session_state
    refs = {r["id"]: r["label"] for r in get_refs()}
    br = lambda series: series.astype(str).str.replace(" | ", "\n", regex=False).str.replace("|", "\n", regex=False)

    df = pd.DataFrame(scenarios)
    df["statut"] = df["id"].map(lambda i: STATUS_LABELS[m["scen_status"].get(i, "pending")][2:].strip())
    df["ref_type"] = df["ref_type"].map(lambda r: refs.get(r, ""))
    for c in ("mode_operatoire", "reco"):
        df[c] = br(df[c])
    cols = {
        "id": "ID", "titre": "Titre", "chemin_id": "Chemin d'attaque", "ref_type": "Scénario de référence",
        "source_risque": "Source de risque", "objectif_vise": "Objectif visé",
        "partie_prenante": "Partie prenante", "valeur_metier": "Valeur métier",
        "evenement_redoute": "Événement redouté", "mode_operatoire": "Mode opératoire",
        "actif": "Bien support ciblé", "vuln": "Vulnérabilité exploitée", "impact_op": "Impact opérationnel",
        "impact_fin": "Impact financier", "impact_reg": "Impact juridique / réglementaire",
        "vraisemblance": "Vraisemblance", "gravite": "Gravité", "score": f"Niveau de risque G×V (/{MAX_SCORE})",
        "reco": "Mesures de sécurité proposées", "statut": "Statut",
    }
    df = df[[c for c in cols if c in df.columns]].rename(columns=cols)

    paths = m["paths"] or []
    pdf = pd.DataFrame(paths)
    if not pdf.empty:
        pdf["statut"] = pdf["id"].map(lambda i: STATUS_LABELS[m["path_status"].get(i, "pending")][2:].strip())
        pdf["ref_type"] = pdf["ref_type"].map(lambda r: refs.get(r, ""))
        pcols = {
            "id": "ID", "source_risque": "Source de risque", "objectif_vise": "Objectif visé",
            "partie_prenante": "Partie prenante", "valeur_metier": "Valeur métier",
            "evenement_redoute": "Événement redouté", "gravite": "Gravité", "chemin": "Chemin d'attaque",
            "ref_type": "Scénario de référence", "justification": "Justification", "statut": "Statut",
        }
        pdf = pdf[[c for c in pcols if c in pdf.columns]].rename(columns=pcols)

    # Empêche l'injection de formules : une cellule commençant par = + - @ serait interprétée par Excel.
    def _safe(v):
        if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
            return "'" + v
        return v

    _map = "map" if hasattr(pd.DataFrame, "map") else "applymap"  # pandas < 2.1 : applymap
    df = getattr(df, _map)(_safe)
    if not pdf.empty:
        pdf = getattr(pdf, _map)(_safe)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Scénarios opérationnels")
        if not pdf.empty:
            pdf.to_excel(writer, index=False, sheet_name="Chemins d'attaque")
        cov = compute_coverage()
        if cov:
            pd.DataFrame(cov).to_excel(writer, index=False, sheet_name="Couverture référentiel")
        for ws in writer.book.worksheets:  # largeur de colonnes lisible
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(
                    60, max(12, max(len(str(c.value or "")) for c in col[:50]) + 2))
    return buf.getvalue()


# ── Barre latérale : mode d'emploi, exemple, nouvelle étude ─────────────

with st.sidebar:
    st.header("🧭 Mode d'emploi")
    st.markdown(
        "1. **Cadrage** : renseignez au minimum le nom et le périmètre (le reste améliore la qualité). "
        "Téléchargez l'exemple ci-dessous pour voir le niveau de détail attendu.\n"
        "2. **Socle** : lisez la synthèse de l'IA et notez les points à clarifier avec le client.\n"
        "3. **Chemins d'attaque** : corrigez si besoin, puis **validez** ou **rejetez** chacun "
        "(en choisissant un motif). Ajoutez ceux que l'IA a oubliés avec « ➕ Ajouter ».\n"
        "4. **Scénarios opérationnels** : relisez, corrigez avec « ✏️ Modifier », validez ou rejetez, "
        "puis **exportez** vers Excel."
    )
    st.divider()
    st.download_button(
        "📄 Télécharger un exemple de cadrage (Word)", data=exemple_word(),
        file_name="Exemple_cadrage_EBIOS_RM.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        width="stretch",
    )
    if st.button("🗑️ Nouvelle étude (tout effacer)", width="stretch"):
        st.session_state["confirm_reset"] = True
    if st.session_state.get("confirm_reset"):
        st.warning("Toutes les données de l'étude en cours seront perdues.")
        c1, c2 = st.columns(2)
        if c1.button("Confirmer", type="primary"):
            reset_state()
            st.rerun()
        if c2.button("Annuler"):
            st.session_state["confirm_reset"] = False
            st.rerun()


# ── En-tête & progression ──────────────────────────────────────────────────

st.title("🛡️ Formind — Outil consultant EBIOS RM")
cols = st.columns(4)
labels = [
    "1. Cadrage et socle", "2. Socle de sécurité",
    "3. Scénarios stratégiques", "4. Scénarios opérationnels",
]
for i, col in enumerate(cols, start=1):
    with col:
        marker = "✅" if st.session_state.step > i else ("▶️" if st.session_state.step == i else "⬜")
        st.markdown(f"**{marker} {labels[i - 1]}**")
st.divider()

if st.session_state.flash:
    st.warning(st.session_state.flash, icon="⚠️")
    st.session_state.flash = ""

# ── Étape 1 : atelier 1, cadrage ───────────────────────────────────────────

if st.session_state.step == 1:
    st.subheader("📋 Atelier 1 — Cadrage et socle de sécurité")
    st.caption("Les champs marqués * sont obligatoires. Plus le cadrage est précis (valeurs métier, biens "
               "supports, événements redoutés, parties prenantes), moins vous aurez de propositions à corriger.")

    c1, c2 = st.columns(2)
    st.session_state.nom = c1.text_input("Nom de l'organisation *", st.session_state.nom)
    st.session_state.taille = c2.selectbox("Taille", TAILLES, index=TAILLES.index(st.session_state.taille))
    st.session_state.secteur = st.selectbox("Secteur d'activité", SECTEURS, index=SECTEURS.index(st.session_state.secteur))
    if st.session_state.secteur == "Autre":
        st.session_state.secteur_autre = st.text_input("Précisez le secteur *", st.session_state.secteur_autre)
    st.session_state.perimetre = st.text_area(
        "Périmètre de l'étude (métier et technique) *",
        st.session_state.perimetre,
        placeholder="ex: activité de paiement et SI Finance, sites de Casablanca et Rabat…",
    )
    st.session_state.valeurs_metier = st.text_area(
        "Valeurs métier",
        st.session_state.valeurs_metier,
        placeholder="ex: Exécution des paiements, gestion des comptes clients, reporting financier réglementaire…",
        help="Processus et informations essentiels à l'organisation.",
    )
    st.session_state.biens_supports = st.text_area(
        "Biens supports",
        st.session_state.biens_supports,
        placeholder="ex: ERP Finance → exécution des paiements ; base de données clients → gestion des comptes ; équipe trésorerie…",
        help="Composants sur lesquels reposent les valeurs métier : systèmes, logiciels, réseaux, personnes, "
             "locaux, prestataires. Indiquez si possible la valeur métier soutenue (bien support → valeur métier).",
    )
    st.session_state.evenements_redoutes = st.text_area(
        "Événements redoutés (avec gravité si connue)",
        st.session_state.evenements_redoutes,
        placeholder="ex: Détournement de virements — Critique ; indisponibilité du SI Finance > 24 h — Grave",
    )

    st.markdown("**Cadrage EBIOS RM**")
    st.session_state.sources_risque = st.multiselect(
        "Sources de risque pressenties (facultatif)",
        SOURCES_RISQUE,
        default=st.session_state.sources_risque,
        help="Typologie indicative inspirée du guide ANSSI. Laissez vide pour laisser le modèle proposer les sources plausibles.",
    )
    st.session_state.scenarios_ref = st.multiselect(
        "Scénarios de référence à couvrir",
        list(SCENARIOS_REFERENCE),
        default=st.session_state.scenarios_ref,
        format_func=lambda rid: SCENARIOS_REFERENCE[rid]["label"],
        help="Chaque scénario coché donnera lieu à au moins un chemin d'attaque proposé, puis à un scénario.",
    )
    st.session_state.scenarios_ref_autre = st.text_area(
        "Autres scénarios de référence (un par ligne)",
        st.session_state.scenarios_ref_autre,
        height=80,
    )

    st.markdown("**Socle de sécurité**")
    st.session_state.reglements = st.multiselect(
        "Référentiels et réglementations applicables",
        REGLEMENTS,
        default=[r for r in st.session_state.reglements if r in REGLEMENTS],
    )
    st.session_state.reglement_autre = st.text_input("Autres obligations spécifiques", st.session_state.reglement_autre)

    c1, c2 = st.columns(2)
    st.session_state.equipe_it = c1.text_input("Effectif équipe IT", st.session_state.equipe_it)
    st.session_state.equipe_sec = c2.text_input("Effectif équipe sécurité", st.session_state.equipe_sec)
    st.session_state.prestataires = st.text_input(
        "Parties prenantes de l'écosystème",
        st.session_state.prestataires,
        placeholder="ex: prestataire d'infogérance, éditeur de l'ERP, partenaires bancaires, filiales, clients…",
    )
    st.session_state.projets_recents = st.text_input("Projets SI en cours ou récents", st.session_state.projets_recents)

    st.session_state.incidents = st.text_area("Incidents de sécurité passés", st.session_state.incidents)
    st.session_state.menaces = st.text_area("État de la menace connu (veille, alertes sectorielles)", st.session_state.menaces)

    st.markdown("**Maturité SSI** (1 = Initial → 5 = Optimisé) — indicateur complémentaire Formind, hors méthode EBIOS RM")
    for k, label in MAT_DIMS:
        st.session_state.maturite[k] = st.slider(label, 1, 5, st.session_state.maturite[k])

    st.session_state.nb_chemins = st.selectbox(
        "Nombre de chemins d'attaque à proposer",
        NB_CHEMINS_OPTIONS,
        index=NB_CHEMINS_OPTIONS.index(st.session_state.nb_chemins),
        help="Un scénario opérationnel sera généré par chemin que vous validerez.",
    )
    st.session_state.notes = st.text_input("Notes du consultant", st.session_state.notes)

    vides = [n for n, v in (("valeurs métier", st.session_state.valeurs_metier),
                            ("biens supports", st.session_state.biens_supports),
                            ("événements redoutés", st.session_state.evenements_redoutes),
                            ("parties prenantes", st.session_state.prestataires)) if not v.strip()]
    if vides:
        st.info("Champs conseillés non renseignés : " + ", ".join(vides)
                + ". L'outil fonctionne sans, mais les propositions seront plus génériques.")

    if st.button("✦ Analyser le socle de sécurité et passer à l'étape 2 →", type="primary", width="stretch"):
        manquants = [n for n, v in (("le nom", st.session_state.nom), ("le périmètre", st.session_state.perimetre)) if not v.strip()]
        if st.session_state.secteur == "Autre" and not st.session_state.secteur_autre.strip():
            manquants.append("le secteur")
        if manquants:
            st.error("⚠ Veuillez renseigner " + ", ".join(manquants) + ".")
        else:
            with st.spinner("Analyse du contexte en cours (environ 20 à 40 secondes)…"):
                try:
                    analysis = call_ia(SYSTEM_SCAN, f"Analyse ce contexte:\n\n{build_context()}")
                    if not isinstance(analysis, dict):
                        raise ValueError("Format de réponse inattendu (objet attendu).")
                    st.session_state.analysis = analysis
                    st.session_state.step = 2
                    st.rerun()
                except Exception as e:
                    erreur_ia(e)

# ── Étape 2 : socle de sécurité ────────────────────────────────────────────

elif st.session_state.step == 2:
    a = st.session_state.analysis or {}
    st.subheader(f"🔍 Atelier 1 — Socle de sécurité — {st.session_state.nom}")

    niveau = a.get("niveau_global", "Non évalué")
    niveau_icon = {"Faible": "🔴", "Moyen": "🟡", "Satisfaisant": "🟢", "Bon": "✅"}.get(niveau, "📊")
    st.info(f"{niveau_icon} **Niveau de maturité SSI estimé : {niveau}**\n\n{a.get('resume', '')}")

    c1, c2, c3 = st.columns(3)
    with c1:
        st.markdown("**✅ Points forts du socle**")
        for f in a.get("forces", []):
            st.markdown(f"- {f}")
    with c2:
        st.markdown("**⚠️ Écarts au socle de sécurité**")
        for g in a.get("gaps", []):
            st.markdown(f"- {g}")
    with c3:
        st.markdown("**🎯 Domaines prioritaires**")
        for p in a.get("priorites", []):
            st.markdown(f"- {p}")

    if a.get("recommandation"):
        st.warning(f"💡 **À clarifier avec le client avant les ateliers 2 et 3**\n\n{a['recommandation']}")
    st.caption("Si le cadrage est incomplet, revenez à l'étape 1 pour le compléter : vous gagnerez du temps de correction ensuite.")

    deja = bool(st.session_state.paths)
    c1, c2, c3 = st.columns([1, 2, 2])
    if c1.button("← Retour"):
        st.session_state.step = 1
        st.rerun()
    if deja and c2.button("Continuer avec les chemins existants →", type="primary", width="stretch"):
        st.session_state.step = 3
        st.rerun()
    label_gen = "↻ Régénérer les chemins (efface chemins et scénarios actuels)" if deja else "🧭 Proposer les scénarios stratégiques →"
    if (c3 if deja else c2).button(label_gen, type="secondary" if deja else "primary", width="stretch"):
        with st.spinner("Construction des couples SR/OV et des chemins d'attaque (environ 30 à 60 secondes)…"):
            try:
                generate_paths()
                st.session_state.step = 3
                st.rerun()
            except Exception as e:
                erreur_ia(e)

# ── Étape 3 : chemins d'attaque (Ateliers 2-3) ─────────────────────────────

elif st.session_state.step == 3:
    paths = st.session_state.paths or []
    pstatus = st.session_state.path_status
    refs = get_refs()
    ref_ids = [""] + [r["id"] for r in refs]
    ref_lbl = {"": "— Aucun (hors référentiel)", **{r["id"]: r["label"] for r in refs}}

    st.subheader(f"🧭 Ateliers 2 et 3 — Scénarios stratégiques — {st.session_state.nom}")
    st.caption(
        "Chaque scénario stratégique associe un couple source de risque / objectif visé (atelier 2) "
        "à un chemin d'attaque passant, le cas échéant, par une partie prenante de l'écosystème (atelier 3). "
        "Corrigez directement les champs si besoin, puis validez ou rejetez : "
        "seuls les chemins validés alimentent l'atelier 4."
    )

    n_val = sum(1 for p in paths if pstatus.get(p["id"]) == "validated")
    n_pend = sum(1 for p in paths if pstatus.get(p["id"], "pending") == "pending")
    m1, m2, m3 = st.columns(3)
    m1.metric("Chemins proposés", len(paths))
    m2.metric("Validés", n_val)
    m3.metric("En attente", n_pend)

    coverage = compute_coverage()
    if coverage:
        st.markdown("**Couverture des scénarios de référence**")
        st.dataframe(coverage, hide_index=True, width="stretch")

    for p in paths:
        pid = p["id"]
        status = pstatus.get(pid, "pending")
        tags = ("  ·  ➕ ajout consultant" if p.get("_origine") == "manuel"
                else ("  ·  ✏️ modifié" if is_modified(p, PATH_EDITABLE) else ""))
        with st.expander(
            f"**{pid}** — {p['source_risque']} → {p['objectif_vise']}  ·  {p['evenement_redoute']}  ·  "
            f"{p['gravite']}  ·  {STATUS_LABELS[status]}{tags}"
        ):
            c1, c2 = st.columns(2)
            p["source_risque"] = choix_ou_saisie(c1, "Source de risque", SOURCES_RISQUE, p["source_risque"], key=f"sr_{pid}")
            p["objectif_vise"] = choix_ou_saisie(c2, "Objectif visé", OBJECTIFS_VISES, p["objectif_vise"], key=f"ov_{pid}")
            p["partie_prenante"] = c1.text_input("Partie prenante de l'écosystème", p["partie_prenante"], key=f"pp_{pid}")
            p["valeur_metier"] = c2.text_input("Valeur métier", p["valeur_metier"], key=f"vm_{pid}")
            p["evenement_redoute"] = c1.text_input("Événement redouté", p["evenement_redoute"], key=f"er_{pid}")
            p["gravite"] = c2.selectbox("Gravité", NIVEAUX_GRAVITE, index=NIVEAUX_GRAVITE.index(p["gravite"]), key=f"gr_{pid}")
            p["chemin"] = st.text_area("Chemin d'attaque", p["chemin"], key=f"ch_{pid}")
            cur_ref = p.get("ref_type", "") if p.get("ref_type", "") in ref_ids else ""
            p["ref_type"] = st.selectbox(
                "Scénario de référence associé", ref_ids, index=ref_ids.index(cur_ref),
                format_func=lambda x: ref_lbl[x], key=f"rf_{pid}",
            )
            if p.get("justification"):
                st.caption(f"Justification : {p['justification']}")
            if p.get("_motif_ajout"):
                st.caption(f"Note du consultant : {p['_motif_ajout']}")
            for msg in warnings_path(p):
                st.warning(msg, icon="⚠️")

            b1, b2, b3, b4 = st.columns([1, 2, 1, 1])
            if b1.button("✓ Valider", key=f"pval_{pid}"):
                pstatus[pid] = "validated"
                p.pop("_motif_rejet", None)
                st.rerun()
            motif = b2.selectbox("Motif en cas de rejet", MOTIFS_REJET, key=f"pmot_{pid}", label_visibility="collapsed")
            if b3.button("✕ Rejeter", key=f"prej_{pid}"):
                pstatus[pid] = "rejected"
                p["_motif_rejet"] = motif
                st.rerun()
            if b4.button("↺ Réinitialiser", key=f"prst_{pid}"):
                pstatus[pid] = "pending"
                p.pop("_motif_rejet", None)
                st.rerun()

    with st.expander("➕ Ajouter un chemin d'attaque manuellement"):
        with st.form("add_path", clear_on_submit=True):
            c1, c2 = st.columns(2)
            f_sr = c1.selectbox("Source de risque", SOURCES_RISQUE + [AUTRE])
            f_ov = c2.selectbox("Objectif visé", OBJECTIFS_VISES + [AUTRE])
            f_sr_libre = c1.text_input("Si « Autre » : précisez la source de risque")
            f_ov_libre = c2.text_input("Si « Autre » : précisez l'objectif visé")
            f_pp = c1.text_input("Partie prenante de l'écosystème", placeholder="ou « Aucune (attaque directe) »")
            f_vm = c2.text_input("Valeur métier")
            f_er = c1.text_input("Événement redouté")
            f_gr = c2.selectbox("Gravité", NIVEAUX_GRAVITE, index=len(NIVEAUX_GRAVITE) - 1)
            f_ch = st.text_area("Chemin d'attaque", placeholder="Source de risque -> partie prenante -> valeur métier -> événement redouté")
            f_rf = st.selectbox("Scénario de référence associé", ref_ids, format_func=lambda x: ref_lbl[x])
            f_mo = st.text_input("Pourquoi ce chemin est important (facultatif)")
            submitted = st.form_submit_button("Ajouter et valider ce chemin")
        if submitted:
            sr = f_sr_libre.strip() if f_sr == AUTRE else f_sr
            ov = f_ov_libre.strip() if f_ov == AUTRE else f_ov
            new = validate_path({
                "source_risque": sr, "objectif_vise": ov, "partie_prenante": f_pp,
                "valeur_metier": f_vm, "evenement_redoute": f_er, "gravite": f_gr,
                "chemin": f_ch, "ref_type": f_rf, "justification": "Ajouté par le consultant.",
            }, {r["id"] for r in refs})
            if new is None:
                st.error("⚠ Tous les champs sont obligatoires (y compris la précision si vous avez choisi « Autre »).")
            else:
                new["source_risque"], new["objectif_vise"] = sr, ov
                new["id"] = f"CA-{len(paths) + 1:02d}"
                new["_origine"] = "manuel"
                new["_motif_ajout"] = f_mo.strip()
                new["_orig"] = {f: new.get(f, "") for f in PATH_EDITABLE}
                paths.append(new)
                pstatus[new["id"]] = "validated"
                st.rerun()

    uncovered = [r["Scénario de référence"] for r in coverage if r["Couverture"].startswith(("🟡", "❌"))]
    if uncovered:
        st.warning(
            "Scénarios de référence sans chemin validé : " + " ; ".join(uncovered)
            + ". Validez un chemin associé ou ajoutez-en un manuellement pour qu'ils soient couverts."
        )
    if n_pend:
        st.info(f"{n_pend} chemin(s) encore en attente : ils ne seront pas traités à l'atelier 4. "
                "Validez-les ou rejetez-les.")

    valid = validated_paths()
    c1, c2 = st.columns([1, 4])
    if c1.button("← Retour"):
        st.session_state.step = 2
        st.rerun()
    if c2.button(
        f"⚡ Générer les scénarios opérationnels ({len(valid)} chemin(s) validé(s)) →",
        type="primary", width="stretch", disabled=not valid,
    ):
        with st.spinner("Génération des scénarios opérationnels EBIOS RM (environ 30 s par lot de 4 chemins)…"):
            try:
                build_scenarios()
                st.session_state.step = 4
                st.rerun()
            except Exception as e:
                erreur_ia(e)

# ── Étape 4 : scénarios opérationnels (Atelier 4) ──────────────────────────

elif st.session_state.step == 4:
    scenarios = st.session_state.scenarios or []
    by_path = {p["id"]: p for p in (st.session_state.paths or [])}
    ref_lbl = {r["id"]: r["label"] for r in get_refs()}
    sstatus = st.session_state.scen_status

    st.subheader(f"🛡 Atelier 4 — Scénarios opérationnels — {st.session_state.nom}")
    st.caption(
        f"{len(scenarios)} scénarios opérationnels, chacun rattaché à un chemin d'attaque validé · "
        "Relisez, corrigez avec « ✏️ Modifier », puis validez ou rejetez chaque scénario. "
        "Le niveau de risque est le produit gravité × vraisemblance."
    )

    actifs_stats = [s for s in scenarios if sstatus.get(s["id"]) != "rejected"]
    crit = sum(1 for s in actifs_stats if s["gravite"] == "Critique")
    val = sum(1 for v in sstatus.values() if v == "validated")
    avg = round(sum(s["score"] for s in actifs_stats) / len(actifs_stats), 1) if actifs_stats else 0
    coverage = compute_coverage()
    covered = sum(1 for r in coverage if r["Couverture"].startswith("✅"))

    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric("Scénarios retenus", len(actifs_stats))
    m2.metric("Gravité critique", crit)
    m3.metric("Validés", val)
    m4.metric(f"Niveau de risque moyen (/{MAX_SCORE})", avg)
    m5.metric("Référentiel couvert", f"{covered}/{len(coverage)}" if coverage else "—")

    if coverage:
        st.markdown("**Couverture des scénarios de référence**")
        st.dataframe(coverage, hide_index=True, width="stretch")

    have = {s["chemin_id"] for s in scenarios}
    todo = [p for p in validated_paths() if p["id"] not in have]
    if todo:
        st.warning(f"{len(todo)} chemin(s) validé(s) n'ont pas de scénario : " + ", ".join(p["id"] for p in todo))
        if st.button("↻ Générer les scénarios manquants"):
            with st.spinner("Génération des scénarios manquants…"):
                try:
                    build_scenarios()
                    st.rerun()
                except Exception as e:
                    erreur_ia(e)

    n_att = sum(1 for s in scenarios if sstatus.get(s["id"], "pending") == "pending")
    c1, c2 = st.columns([1, 4])
    if c1.button("← Retour"):
        st.session_state.step = 3
        st.rerun()
    if scenarios:
        c2.download_button(
            "↓ Exporter Excel (scénarios, chemins, couverture)",
            data=export_excel(scenarios),
            file_name=f"EBIOS_RM_{st.session_state.nom.replace(' ', '_')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )
        if n_att:
            c2.caption(f"{n_att} scénario(s) encore en attente : ils seront exportés avec le statut « En attente ».")

    filtre = st.radio("Filtrer", ["Tous", "Gravité critique", "Validés", "En attente", "Rejetés"], horizontal=True)
    shown = scenarios
    if filtre == "Gravité critique":
        shown = [s for s in scenarios if s["gravite"] == "Critique"]
    elif filtre == "Validés":
        shown = [s for s in scenarios if sstatus.get(s["id"]) == "validated"]
    elif filtre == "En attente":
        shown = [s for s in scenarios if sstatus.get(s["id"], "pending") == "pending"]
    elif filtre == "Rejetés":
        shown = [s for s in scenarios if sstatus.get(s["id"]) == "rejected"]

    for s in shown:
        sid = s["id"]
        status = sstatus.get(sid, "pending")
        tag = "  ·  ✏️ modifié" if is_modified(s, SCEN_EDITABLE) else ""
        with st.expander(
            f"**{sid}** — {s['titre']}  ·  {s['chemin_id']}  ·  V : {s['vraisemblance']} / G : {s['gravite']}  ·  "
            f"risque {s['score']}/{MAX_SCORE}  ·  {STATUS_LABELS[status]}{tag}"
        ):
            path = by_path.get(s["chemin_id"], {})
            edit = st.toggle("✏️ Modifier ce scénario", key=f"edit_{sid}")
            c1, c2 = st.columns(2)
            c1.markdown(f"**Source de risque**\n\n{s['source_risque']} → *{s['objectif_vise']}*")
            c1.markdown(f"**Chemin d'attaque {s['chemin_id']}**\n\n{path.get('chemin', '—')}")
            c1.markdown(f"**Partie prenante de l'écosystème**\n\n{s['partie_prenante']}")
            c2.markdown(f"**Valeur métier → événement redouté**\n\n{s['valeur_metier']} → {s['evenement_redoute']}")
            c2.markdown(f"**Gravité (héritée du chemin)** : {s['gravite']}")

            if edit:
                st.caption("Source, objectif, partie prenante, valeur métier, événement redouté et gravité "
                           "se corrigent à l'étape 3 (chemin d'attaque).")
                s["titre"] = st.text_input("Titre", s["titre"], key=f"ti_{sid}")
                e1, e2 = st.columns(2)
                s["actif"] = e1.text_input("Bien support ciblé", s["actif"], key=f"ac_{sid}")
                s["vuln"] = e2.text_input("Vulnérabilité exploitée", s["vuln"], key=f"vu_{sid}")
                s["vraisemblance"] = e1.selectbox(
                    "Vraisemblance", NIVEAUX_VRAISEMBLANCE,
                    index=NIVEAUX_VRAISEMBLANCE.index(s["vraisemblance"]), key=f"vr_{sid}",
                )
                s["score"] = score_scenario(s["vraisemblance"], s["gravite"])
                e2.metric("Niveau de risque recalculé", f"{s['score']}/{MAX_SCORE}")
                s["mode_operatoire"] = _pipe(st.text_area(
                    "Mode opératoire (une phase par ligne)", s["mode_operatoire"].replace(" | ", "\n"),
                    key=f"mo_{sid}", height=140,
                ))
                s["impact_op"] = st.text_area("Impact opérationnel", s["impact_op"], key=f"io_{sid}")
                s["impact_fin"] = st.text_area("Impact financier", s["impact_fin"], key=f"if_{sid}")
                s["impact_reg"] = st.text_area("Impact juridique / réglementaire", s["impact_reg"], key=f"ir_{sid}")
                s["reco"] = _pipe(st.text_area(
                    "Mesures de sécurité (une par ligne)", s["reco"].replace(" | ", "\n"), key=f"re_{sid}", height=120,
                ))
                if is_modified(s, SCEN_EDITABLE) and st.button("↶ Revenir à la proposition de l'IA", key=f"undo_{sid}"):
                    for f in SCEN_EDITABLE:
                        s[f] = s["_orig"][f]
                    s["score"] = score_scenario(s["vraisemblance"], s["gravite"])
                    for wk in ("ti", "ac", "vu", "vr", "mo", "io", "if", "ir", "re"):
                        st.session_state.pop(f"{wk}_{sid}", None)
                    st.rerun()
            else:
                c1.markdown(f"**Bien support ciblé**\n\n{s['actif']}")
                c1.markdown(f"**Vulnérabilité exploitée**\n\n{s['vuln']}")
                c2.markdown(f"**Impact opérationnel**\n\n{s['impact_op']}")
                c2.markdown(f"**Impact financier**\n\n{s['impact_fin']}")
                c2.markdown(f"**Impact juridique / réglementaire**\n\n{s['impact_reg']}")
                st.markdown("**Mode opératoire**\n\n" + "\n".join(
                    f"{i}. {step.strip()}" for i, step in enumerate(s["mode_operatoire"].split("|"), start=1)
                ))
                st.info("↳ **Mesures de sécurité proposées** (à reprendre dans le traitement du risque, atelier 5)\n\n"
                        + "\n\n".join(f"- {r.strip()}" for r in s["reco"].split("|")))
            if s.get("ref_type") in ref_lbl:
                st.caption(f"Scénario de référence : {ref_lbl[s['ref_type']]}")
            for msg in warnings_scenario(s):
                st.warning(msg, icon="⚠️")

            b1, b2, b3, b4 = st.columns([1, 2, 1, 1])
            if b1.button("✓ Valider", key=f"val_{sid}"):
                sstatus[sid] = "validated"
                s.pop("_motif_rejet", None)
                st.rerun()
            motif = b2.selectbox("Motif en cas de rejet", MOTIFS_REJET, key=f"smot_{sid}", label_visibility="collapsed")
            if b3.button("✕ Rejeter", key=f"rej_{sid}"):
                sstatus[sid] = "rejected"
                s["_motif_rejet"] = motif
                st.rerun()
            if b4.button("↺ Réinitialiser", key=f"rst_{sid}"):
                sstatus[sid] = "pending"
                s.pop("_motif_rejet", None)
                st.rerun()


# ── Enregistrement silencieux des indicateurs (fin de chaque exécution) ─────
if st.session_state.paths:
    save_indicators()